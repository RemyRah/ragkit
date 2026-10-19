"""Load documents from disk and split them into retrievable chunks.

Chunking respects Markdown structure: it breaks on headings first, then packs
paragraphs up to a target size with a small overlap so a fact that straddles a
boundary still lands whole in at least one chunk.
"""

from __future__ import annotations

import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

# Files we know how to read as plain text.
TEXT_SUFFIXES = {".md", ".markdown", ".txt", ".rst"}


class MissingExtra(Exception):
    """Raised by a reader whose optional dependency isn't installed."""

    def __init__(self, package: str):
        super().__init__(package)
        self.package = package


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise MissingExtra("pypdf") from None
    return "\n\n".join((page.extract_text() or "") for page in PdfReader(path).pages)


def _read_docx(path: Path) -> str:
    try:
        import docx
    except ImportError:
        raise MissingExtra("python-docx") from None
    lines = []
    for para in docx.Document(path).paragraphs:
        style = para.style.name if para.style is not None else ""
        # Keep Word headings as Markdown headings so chunking/citations use them.
        if style.startswith("Heading") and style[-1:].isdigit():
            lines.append("#" * int(style[-1]) + " " + para.text)
        else:
            lines.append(para.text)
    return "\n\n".join(lines)


def _missing_reader(package: str):
    """A reader that always reports `package` as missing (handy for tests)."""

    def read(path: Path) -> str:
        raise MissingExtra(package)

    return read


# Binary formats, each behind an optional extra.
READERS = {".pdf": _read_pdf, ".docx": _read_docx}


@dataclass
class Document:
    """One source file loaded from disk."""

    source: str
    text: str


@dataclass
class Chunk:
    """A retrievable slice of a document, with a pointer back to its source."""

    source: str
    heading: str
    text: str
    index: int
    metadata: dict = field(default_factory=dict)

    @property
    def citation(self) -> str:
        loc = self.heading or f"chunk {self.index}"
        return f"{self.source} — {loc}"


def load_docs(path: str | Path) -> list[Document]:
    """Load every supported file under `path` (file or directory).

    Text/Markdown always load; PDF and DOCX load when their optional extras
    (pypdf, python-docx) are installed, otherwise they are skipped with a warning.
    """

    root = Path(path)
    if not root.exists():
        raise FileNotFoundError(f"No such path: {root}")

    files: list[Path]
    if root.is_file():
        files = [root]
    else:
        files = sorted(
            p
            for p in root.rglob("*")
            if p.is_file() and (p.suffix.lower() in TEXT_SUFFIXES or p.suffix.lower() in READERS)
        )

    docs: list[Document] = []
    warned: set[str] = set()
    for f in files:
        try:
            reader = READERS.get(f.suffix.lower())
            text = reader(f) if reader else f.read_text(encoding="utf-8")
        except MissingExtra as e:
            if e.package not in warned:
                warned.add(e.package)
                print(
                    f"ragkit: skipping {f.suffix} files — run `pip install {e.package}` to read them.",
                    file=sys.stderr,
                )
            continue
        except UnicodeDecodeError:
            # Skip binaries that slipped through; a real deployment would log this.
            continue
        if text.strip():
            docs.append(Document(source=str(f), text=text))
    return docs


_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")


def _split_sections(text: str) -> list[tuple[str, str]]:
    """Split Markdown into (heading, body) sections. Heading is '' before the first one."""

    sections: list[tuple[str, list[str]]] = [("", [])]
    for line in text.splitlines():
        m = _HEADING_RE.match(line)
        if m:
            sections.append((m.group(2).strip(), []))
        else:
            sections[-1][1].append(line)
    return [(h, "\n".join(body).strip()) for h, body in sections if "\n".join(body).strip()]


def chunk_documents(
    docs: list[Document],
    target_chars: int = 900,
    overlap_chars: int = 150,
) -> list[Chunk]:
    """Turn documents into overlapping chunks packed to ~`target_chars`."""

    chunks: list[Chunk] = []
    idx = 0
    for doc in docs:
        for heading, body in _split_sections(doc.text):
            for piece in _pack(body, target_chars, overlap_chars):
                chunks.append(
                    Chunk(source=doc.source, heading=heading, text=piece, index=idx)
                )
                idx += 1
    return chunks


def _pack(body: str, target: int, overlap: int) -> list[str]:
    """Greedily pack paragraphs up to `target`, carrying `overlap` chars forward."""

    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", body) if p.strip()]
    out: list[str] = []
    buf = ""
    for para in paragraphs:
        if buf and len(buf) + len(para) + 2 > target:
            out.append(buf)
            buf = buf[-overlap:] + "\n\n" + para if overlap else para
        else:
            buf = f"{buf}\n\n{para}" if buf else para
    if buf.strip():
        out.append(buf)
    return out
